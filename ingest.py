"""
ingest.py — загрузка материалов в базу знаний.

Поддерживаемые форматы:
  - PDF (.pdf)
  - Word (.docx)
  - Таблицы (.xlsx, .csv)
  - Текст (.txt)
  - YouTube (ссылка на видео)

Использование:
  python3 ingest.py                        # обработать все файлы из knowledge_base/raw/
  python3 ingest.py https://youtu.be/...   # добавить YouTube видео
  python3 ingest.py путь/к/файлу.pdf       # добавить конкретный файл
"""

import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from knowledge_base.kb_store import add_chunks, count

RAW_DIR = Path("knowledge_base/raw")


def chunk_text(text: str, chunk_size: int = 600, overlap: int = 80) -> list[str]:
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i : i + chunk_size])
        chunks.append(chunk)
        i += chunk_size - overlap
    return [c for c in chunks if len(c.strip()) > 60]


def ingest_pdf(path: Path):
    import PyPDF2
    print(f"PDF: {path.name}")
    text = ""
    with open(path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"
    if not text.strip():
        print("  ! Текст не извлечён (возможно, PDF с картинками)")
        return
    add_chunks(chunk_text(text), path.name, "pdf")


def ingest_docx(path: Path):
    import docx
    print(f"DOCX: {path.name}")
    doc = docx.Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    add_chunks(chunk_text(text), path.name, "docx")


def ingest_excel(path: Path):
    print(f"Excel/CSV: {path.name}")
    if path.suffix == ".csv":
        text = path.read_text(encoding="utf-8", errors="ignore")
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                row_text = "\t".join(str(c) for c in row if c is not None)
                if row_text.strip():
                    rows.append(row_text)
        text = "\n".join(rows)
    add_chunks(chunk_text(text, chunk_size=300), path.name, "table")


def ingest_youtube(url: str):
    from youtube_transcript_api import YouTubeTranscriptApi
    import re
    print(f"YouTube: {url}")

    match = re.search(r"(?:v=|youtu\.be/)([a-zA-Z0-9_-]{11})", url)
    if not match:
        print("  ! Не удалось извлечь ID видео из ссылки")
        return

    video_id = match.group(1)
    try:
        transcript = YouTubeTranscriptApi.get_transcript(video_id, languages=["ru", "en"])
        text = " ".join(entry["text"] for entry in transcript)
        add_chunks(chunk_text(text), url, "youtube")
    except Exception as e:
        print(f"  ! Ошибка получения субтитров: {e}")
        print("  Совет: у видео должны быть субтитры (авто или ручные)")


def ingest_file(path: Path):
    suffix = path.suffix.lower()
    handlers = {
        ".pdf": ingest_pdf,
        ".docx": ingest_docx,
        ".doc": ingest_docx,
        ".xlsx": ingest_excel,
        ".xls": ingest_excel,
        ".csv": ingest_excel,
        ".txt": lambda p: add_chunks(
            chunk_text(p.read_text(encoding="utf-8", errors="ignore")), p.name, "txt"
        ),
    }
    handler = handlers.get(suffix)
    if handler:
        handler(path)
    else:
        print(f"  ? Формат не поддерживается: {path.name}")


def ingest_all_raw():
    files = [f for f in RAW_DIR.iterdir() if f.is_file()]
    if not files:
        print(f"Папка {RAW_DIR}/ пуста.")
        print("Добавь туда файлы (PDF, DOCX, XLSX, TXT) и запусти снова.")
        return
    for f in sorted(files):
        ingest_file(f)


if __name__ == "__main__":
    if len(sys.argv) > 1:
        arg = sys.argv[1]
        if arg.startswith("http"):
            ingest_youtube(arg)
        else:
            ingest_file(Path(arg))
    else:
        ingest_all_raw()

    print(f"\nИтого в базе знаний: {count()} фрагментов")
